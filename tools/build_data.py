#!/usr/bin/env python3
"""
Sinh file public/data.js từ lịch thi đấu (PDF hoặc Excel).

Cách dùng:
    python3 tools/build_data.py "Lich thi dau.pdf"           # cần pdftotext (poppler)
    python3 tools/build_data.py "Lich thi dau.xlsx"          # cần openpyxl
    python3 tools/build_data.py file.pdf --region "Hải Phòng" --date 13/09/2026

Cột mong đợi (Excel): STT | Thời gian | Sân | Nội dung | Bảng | Đội 1 | Thành viên 1 | Thành viên 2 | Đội 2
Script tự sinh thêm các cặp đấu vòng tròn còn thiếu (nếu lịch chưa liệt kê đủ),
các trận này để trống giờ/sân và hiện dấu "—" trên web.
"""
import sys, os, re, json, itertools, argparse, subprocess, collections

ROW_RE = re.compile(
    r'\s*(\d+)\s+(\d\d:\d\d\s*-\s*\d\d:\d\d)\s+(Sân\s*\d+)\s+'
    r'(ĐÔI NAM NỮ|ĐÔI NAM|ĐÔI NỮ)\s+([A-Z])\s+(\S+)\s+(.+?)\s{2,}(.+?)\s{2,}(\S+)\s*$')

CAT_NORM = {"ĐÔI NAM": "ĐÔI NAM", "ĐÔI NAM NỮ": "ĐÔI NAM NỮ", "ĐÔI NỮ": "ĐÔI NỮ"}

# Dòng vòng loại trực tiếp: ... | Play-off / Bán kết / Chung kết | nhãn đội 1 | nhãn đội 2
KO_RE = re.compile(
    r'\s*(\d+)\s+(\d\d:\d\d\s*-\s*\d\d:\d\d)\s+(Sân\s*\d+)\s+'
    r'(ĐÔI NAM NỮ|ĐÔI NAM|ĐÔI NỮ)\s+(Play-?off|Bán kết|Chung kết)\s+(.+?)\s{2,}(.+?)\s*$')
ROUND_NORM = {"play-off": "Play-off", "playoff": "Play-off",
              "bán kết": "Bán kết", "chung kết": "Chung kết"}


def parse_slot(label):
    """Chuyển nhãn ('Nhất A', 'Nhì XSN', 'Thắng Play-off', 'Thắng BK1') -> slot."""
    t = re.sub(r'\s+', ' ', label.strip())
    m = re.match(r'^Nhất\s+([A-Z])$', t, re.I)
    if m:
        return {"t": "w", "grp": m.group(1).upper(), "label": t}
    if re.match(r'^Nhì\s*(XSN|XS|xuất sắc.*)$', t, re.I):
        return {"t": "bs", "label": "Nhì XSN"}
    if re.match(r'^Thắng\s*Play-?off$', t, re.I):
        return {"t": "m", "ref": "PO1", "label": t}
    m = re.match(r'^Thắng\s*BK\s*(\d)$', t, re.I)
    if m:
        return {"t": "m", "ref": "BK" + m.group(1), "label": t}
    m = re.match(r'^Nhì\s+([A-Z])$', t, re.I)
    if m:
        return {"t": "r2", "grp": m.group(1).upper(), "label": t}
    return {"t": "?", "label": t}


def from_pdf(path):
    txt = subprocess.run(["pdftotext", "-layout", path, "-"],
                         capture_output=True, text=True, check=True).stdout
    rows, ko = [], []
    for line in txt.splitlines():
        m = ROW_RE.match(line)
        if m:
            rows.append(dict(no=int(m.group(1)), time=re.sub(r'\s*-\s*', ' - ', m.group(2)),
                             court=m.group(3).replace("  ", " "), cat=m.group(4), grp=m.group(5),
                             t1=m.group(6), n1=m.group(7).strip(),
                             n2=m.group(8).strip(), t2=m.group(9)))
            continue
        k = KO_RE.match(line)
        if k:
            ko.append(dict(no=int(k.group(1)), time=re.sub(r'\s*-\s*', ' - ', k.group(2)),
                           court=k.group(3).replace("  ", " "), cat=k.group(4),
                           rnd=ROUND_NORM.get(k.group(5).lower(), k.group(5)),
                           l1=k.group(6), l2=k.group(7)))
    return rows, ko


def from_xlsx(path):  # noqa: chưa hỗ trợ dòng KO -> trả list rỗng
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            c = [("" if x is None else str(x).strip()) for x in r]
            if len(c) < 9:
                continue
            # tìm cột nội dung
            try:
                ci = next(i for i, x in enumerate(c) if x.upper() in CAT_NORM)
            except StopIteration:
                continue
            if not re.match(r'^\d+$', c[0]):
                continue
            rows.append(dict(no=int(c[0]), time=c[1], court=c[2], cat=c[ci].upper(),
                             grp=c[ci + 1], t1=c[ci + 2], n1=c[ci + 3],
                             n2=c[ci + 4], t2=c[ci + 5]))
    return rows, []


def build(rows, ko_rows, meta):
    teams, groups = {}, collections.defaultdict(list)
    for r in rows:
        for t, n in ((r['t1'], r['n1']), (r['t2'], r['n2'])):
            if t not in teams:
                teams[t] = {'code': t, 'players': n, 'cat': r['cat'], 'grp': r['grp']}
                groups[(r['cat'], r['grp'])].append(t)
    matches, seen = [], set()
    for r in rows:
        seen.add((r['cat'], r['grp'], frozenset([r['t1'], r['t2']])))
        matches.append({'id': f"M{r['no']:03d}", 'no': r['no'], 'time': r['time'],
                        'court': r['court'], 'cat': r['cat'], 'grp': r['grp'],
                        't1': r['t1'], 't2': r['t2'], 'scheduled': True})
    nid = 1000
    for (cat, grp), ts in groups.items():
        for a, b in itertools.combinations(sorted(ts), 2):
            if (cat, grp, frozenset([a, b])) not in seen:
                nid += 1
                matches.append({'id': f"X{nid}", 'no': None, 'time': '', 'court': '',
                                'cat': cat, 'grp': grp, 't1': a, 't2': b, 'scheduled': False})
    matches.sort(key=lambda m: (m['no'] is None, m['no'] or 0, m['cat'], m['grp']))

    # ---- vòng loại trực tiếp ----
    ORDER = {"Play-off": 0, "Bán kết": 1, "Chung kết": 2}
    ko_rows = sorted(ko_rows, key=lambda k: (ORDER.get(k['rnd'], 9), k['no']))
    cnt = collections.defaultdict(int)
    ko = []
    for k in ko_rows:
        pre = {"Play-off": "PO", "Bán kết": "BK", "Chung kết": "CK"}.get(k['rnd'], "KO")
        cnt[(k['cat'], pre)] += 1
        sid = f"{pre}{cnt[(k['cat'], pre)]}"
        ko.append({'id': f"{k['cat']}|{sid}", 'sid': sid, 'no': k['no'], 'time': k['time'],
                   'court': k['court'], 'cat': k['cat'], 'round': k['rnd'],
                   's1': parse_slot(k['l1']), 's2': parse_slot(k['l2'])})
    return {'meta': meta, 'teams': teams, 'matches': matches, 'knockout': ko,
            'groups': {f"{c}|{g}": sorted(ts) for (c, g), ts in groups.items()}}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--region", default="Hà Nội")
    ap.add_argument("--date", default="12/09/2026")
    ap.add_argument("-o", "--out", default=os.path.join(os.path.dirname(__file__), "..", "data.js"))
    a = ap.parse_args()

    rows, ko_rows = from_xlsx(a.input) if a.input.lower().endswith((".xlsx", ".xlsm")) else from_pdf(a.input)
    if not rows:
        sys.exit("Không đọc được dòng lịch nào — kiểm tra lại định dạng file.")
    meta = {'region': a.region, 'date': a.date,
            'title': f"Giải Cầu lông Nội bộ An Phát Holdings 2026",
            'shortTitle': "Cầu lông Nội bộ APH 2026"}
    data = build(rows, ko_rows, meta)
    with open(a.out, "w", encoding="utf-8") as f:
        f.write("// Dữ liệu giải — SINH TỰ ĐỘNG bởi tools/build_data.py. Đừng sửa tay.\n"
                "window.TOURNAMENT_DATA = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n")
    gen = sum(1 for m in data['matches'] if not m['scheduled'])
    print(f"OK -> {a.out}\n  {len(data['teams'])} đội, {len(data['matches'])} trận vòng bảng "
          f"({len(rows)} từ file + {gen} tự sinh vòng tròn), "
          f"{len(data['knockout'])} trận loại trực tiếp")
    for k in data['knockout']:
        print(f"    #{k['no']:>3} {k['cat']:<12} {k['round']:<10} {k['s1']['label']} vs {k['s2']['label']}")
    for k in sorted(data['groups']):
        print("   ", k, len(data['groups'][k]), "đội")


if __name__ == "__main__":
    main()
